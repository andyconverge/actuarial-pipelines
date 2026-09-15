#!/usr/bin/env python3
"""
Load the Settlement tab of the KSKJ monthly report into converge-database.kskj.settlement.

Companion to KSKJ etl.ipynb, following the same conventions: set_month as the period key,
service-account auth, lowercase snake_case columns.

The Settlement sheet is a fixed-layout statement. Rows are line items in five sections;
columns are two blocks -- Gross (E:M) and Net at quota share (O:W) -- each with a Total plus
eight cohorts (MYG03-MYG10). This reshapes that grid into a long, narrow fact table,
validates it against itself and against the seriatim/transactional sheets, then loads it.

Two deliberate departures from KSKJ etl.ipynb:

  1. pandas_gbq.to_gbq(df, ...) rather than df.to_gbq(...). The DataFrame method is
     deprecated as of pandas 2.2.0 and is removed in pandas 3.0.
  2. Delete + insert on set_month rather than a bare if_exists='append', so re-running a
     month is idempotent instead of silently doubling it.

Use as a module -- the drop-in for KSKJ etl.ipynb
-------------------------------------------------
    sys.path.append('../actuarial-pipelines/settlement')
    from settlement_etl import load_settlement

    # credentials accepts the CREDS path directly, or an already-built Credentials object
    load_settlement(file, set_month=set_month, credentials=CREDS, dry_run=False)

Use from the command line
-------------------------
    # extract and validate only, no writes
    python settlement_etl.py "I:/.../KSKJ Converge Report 20260731.xlsx" --dry-run

    # load one month
    python settlement_etl.py "I:/.../KSKJ Converge Report 20260731.xlsx" \
        --creds ../converge-database-0331482f2ee5.json

    # back-load history
    python settlement_etl.py "I:/New Structure/Actuarial New/Database/KSKJ/*/KSKJ Converge Report *.xlsx" \
        --creds ../converge-database-0331482f2ee5.json

    # dump to CSV instead of loading
    python settlement_etl.py report.xlsx --dry-run --csv settlement.csv

Requires: openpyxl, pandas, pandas-gbq, google-cloud-bigquery, db-dtypes
(the GCP packages are imported lazily, so --dry-run works without them).
"""

from __future__ import annotations

import argparse
import glob
import logging
import os
import re
import sys
import warnings
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
import pandas as pd

log = logging.getLogger("settlement")

# ---------------------------------------------------------------------------
# Destination -- matching KSKJ etl.ipynb
# ---------------------------------------------------------------------------

PROJECT = "converge-database"
DATASET = "kskj"
TABLE = "settlement"
CREDS = "../converge-database-0331482f2ee5.json"

CEDENT = "KSKJ"
TREATY = "MYGA"

# ---------------------------------------------------------------------------
# Sheet layout
#
# These constants are the only thing that needs changing if KSKJ restructures the
# template. Two quirks of the source file:
#
#   * The sheet reports its dimension as A1:XFB135 -- roughly 16,000 columns of junk
#     metadata. Every read is bounded to the real data range rather than trusting
#     ws.max_column.
#   * Values are formula results, so the workbook opens with data_only=True. That reads
#     Excel's cached result, which only exists if the file was last saved by Excel.
#     Amounts coming back as None means the cache is empty -- open it in Excel and re-save.
# ---------------------------------------------------------------------------

SHEET_NAME = "Settlement"
PERIOD_CELL = "F4"

COL_LINE_CODE = 1  # A -- 'A'..'K' settlement step labels
COL_LINE_ITEM = 2  # B -- line item description
COL_LINE_NOTE = 3  # C -- supplementary description on some rows

QS_ROW = 6         # quota share applied to each column
HEADER_ROW = 7     # 'Total', 'MYG03'..'MYG10'
FIRST_DATA_ROW = 8
LAST_DATA_ROW = 200  # generous upper bound; blank rows are skipped

BLOCKS = {"gross": range(5, 14), "net": range(15, 24)}

DUST_THRESHOLD = 1e-6  # float noise below this is stored as exact 0
AMOUNT_PRECISION = 9

# BigQuery NUMERIC is DECIMAL(38, 9). Amounts are quantized to this scale on load.
NUMERIC_SCALE = 9
NUMERIC_QUANTUM = Decimal(1).scaleb(-NUMERIC_SCALE)  # Decimal('1E-9')

ROLLFORWARD_SECTIONS = (4, 5)
ROLLFORWARD_LABELS = {
    "begin": "Values at the Beginning",
    "change": "Change in Values",
    "end": "Values at the End",
}

SECTION_RE = re.compile(r"^Section\s+(\d+)")

# ---------------------------------------------------------------------------
# BigQuery schema
#
# amount is NUMERIC rather than the FLOAT64 the other KSKJ tables use. Deliberate: this
# table exists to prove equality to the cent, and binary floats cannot represent cents
# exactly -- the whole reason the tie-outs need a tolerance. NUMERIC is exact decimal, so
# settlement figures compare and sum without drift. Joins on set_month are unaffected.
# ---------------------------------------------------------------------------

BQ_SCHEMA = [
    {"name": "set_month", "type": "STRING", "mode": "REQUIRED"},
    {"name": "report_date", "type": "DATE", "mode": "REQUIRED"},
    {"name": "cedent", "type": "STRING", "mode": "REQUIRED"},
    {"name": "treaty", "type": "STRING", "mode": "NULLABLE"},
    {"name": "section", "type": "INT64", "mode": "NULLABLE"},
    {"name": "section_label", "type": "STRING", "mode": "NULLABLE"},
    {"name": "line_code", "type": "STRING", "mode": "NULLABLE"},
    {"name": "line_item", "type": "STRING", "mode": "REQUIRED"},
    {"name": "line_note", "type": "STRING", "mode": "NULLABLE"},
    {"name": "row_ord", "type": "INT64", "mode": "REQUIRED"},
    {"name": "basis", "type": "STRING", "mode": "REQUIRED"},
    {"name": "cohort", "type": "STRING", "mode": "REQUIRED"},
    {"name": "quota_share", "type": "FLOAT64", "mode": "NULLABLE"},
    {"name": "amount", "type": "NUMERIC", "mode": "NULLABLE"},
    {"name": "source_file", "type": "STRING", "mode": "NULLABLE"},
    {"name": "loaded_at", "type": "TIMESTAMP", "mode": "NULLABLE"},
]

COLUMN_ORDER = [field["name"] for field in BQ_SCHEMA]

DDL_TEMPLATE = """
CREATE TABLE IF NOT EXISTS `{fqn}` (
  set_month     STRING    NOT NULL,
  report_date   DATE      NOT NULL,
  cedent        STRING    NOT NULL,
  treaty        STRING,
  section       INT64,
  section_label STRING,
  line_code     STRING,
  line_item     STRING    NOT NULL,
  line_note     STRING,
  row_ord       INT64     NOT NULL,
  basis         STRING    NOT NULL,
  cohort        STRING    NOT NULL,
  quota_share   FLOAT64,
  amount        NUMERIC,
  source_file   STRING,
  loaded_at     TIMESTAMP
)
PARTITION BY report_date
CLUSTER BY set_month, basis, cohort, line_code
"""

DELETE_TEMPLATE = """
DELETE FROM `{fqn}`
WHERE set_month = @set_month
  AND cedent = @cedent
"""


class ValidationError(AssertionError):
    """Raised when the settlement statement fails its checks."""


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------


def clean_text(value) -> str | None:
    """Normalise a label cell to a single-spaced string, or None if empty."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()
    return text or None


def clean_amount(value) -> float | None:
    """Return a float for numeric cells, None otherwise. Squashes float dust to exact 0."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    amount = float(value)
    if abs(amount) < DUST_THRESHOLD:  # e.g. 'Other Increases' = -9.7e-08
        return 0.0
    return round(amount, AMOUNT_PRECISION)


def read_period(worksheet) -> date:
    """Read the accounting period end date from the header."""
    value = worksheet[PERIOD_CELL].value
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        raise ValueError(f"{PERIOD_CELL} is not a date: {value!r} -- template may have shifted")
    return value


def read_column_map(worksheet) -> list[tuple[str, int, str, float | None]]:
    """Return (basis, column_index, cohort, quota_share) for every data column."""
    columns = []
    for basis, col_range in BLOCKS.items():
        for col in col_range:
            cohort = clean_text(worksheet.cell(HEADER_ROW, col).value)
            if not cohort:
                continue
            if cohort.lower() == "total":
                cohort = "TOTAL"
            columns.append((basis, col, cohort, clean_amount(worksheet.cell(QS_ROW, col).value)))
    if not columns:
        raise ValueError(f"no cohort headers on row {HEADER_ROW} -- template may have shifted")
    return columns


def extract_settlement(path: str | Path, cedent: str = CEDENT, treaty: str = TREATY) -> pd.DataFrame:
    """Read the Settlement sheet into one row per (line item, basis, cohort).

    Two rules keep this robust to the template's quirks.

    Key on row position, not label. The sheet contains unlabelled 'Blanks'/'Blank'
    placeholder rows that KSKJ may populate later, and Sections 4 and 5 reuse identical
    labels ('Values at the Beginning of the Month'). row_ord plus section identifies a
    line; the label is descriptive.

    Emit cell by cell. Some lines exist on only one basis -- E (Monthly Expenses owed from
    Reinsurer) and F through K are net-only. Emitting per cell handles that without
    inventing null gross figures.

    Zeros are kept. Skipping them would make month-over-month diffs unreliable, since a
    line dropping to zero is a real event.
    """
    path = Path(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"{path.name}: no '{SHEET_NAME}' sheet (found {workbook.sheetnames})")
    worksheet = workbook[SHEET_NAME]

    report_date = read_period(worksheet)
    set_month = report_date.strftime("%Y%m")  # same period key as the rest of the pipeline
    columns = read_column_map(worksheet)

    records = []
    section = section_label = None

    for row in range(FIRST_DATA_ROW - 1, LAST_DATA_ROW + 1):
        col_a = clean_text(worksheet.cell(row, COL_LINE_CODE).value)

        # Section headers live in column A and reset the current section.
        match = SECTION_RE.match(col_a) if col_a else None
        if match:
            section, section_label = int(match.group(1)), col_a
            continue
        if row < FIRST_DATA_ROW:
            continue

        line_item = clean_text(worksheet.cell(row, COL_LINE_ITEM).value)
        if not line_item:
            continue  # unlabelled filler rows carry no meaning

        # Short column-A values are settlement step codes ('A'..'K').
        line_code = col_a if col_a and len(col_a) <= 3 else None
        line_note = clean_text(worksheet.cell(row, COL_LINE_NOTE).value)

        for basis, col, cohort, quota_share in columns:
            amount = clean_amount(worksheet.cell(row, col).value)
            if amount is None:
                continue
            records.append(
                {
                    "set_month": set_month,
                    "report_date": report_date,
                    "cedent": cedent,
                    "treaty": treaty,
                    "section": section,
                    "section_label": section_label,
                    "line_code": line_code,
                    "line_item": line_item,
                    "line_note": line_note,
                    "row_ord": row,
                    "basis": basis,
                    "cohort": cohort,
                    "quota_share": quota_share,
                    "amount": amount,
                    "source_file": path.name,
                }
            )

    frame = pd.DataFrame.from_records(records)
    if frame.empty:
        raise ValueError(f"{path.name}: extracted zero rows -- check the layout constants")

    key = ["set_month", "row_ord", "basis", "cohort"]
    duplicates = int(frame.duplicated(key).sum())
    if duplicates:
        raise ValueError(f"{path.name}: {duplicates} duplicate rows on {key}")

    frame["loaded_at"] = pd.Timestamp.utcnow()

    log.info(
        "%s: %d rows, set_month %s, %d line items, cohorts %s",
        path.name, len(frame), set_month, frame.row_ord.nunique(), sorted(frame.cohort.unique()),
    )
    return frame


# ---------------------------------------------------------------------------
# Transactional sheets, for cross-checking
# ---------------------------------------------------------------------------


def normalise_columns(frame: pd.DataFrame) -> pd.DataFrame:
    """Header normalisation matching KSKJ etl.ipynb, plus dropping unnamed columns."""
    frame = frame.copy()
    frame.columns = (
        frame.columns.str.strip()
        .map(str.lower)
        .map(lambda c: c.replace(" ", "_"))
        .map(lambda c: c.replace("/", "_"))
        .map(lambda c: c.replace("+", "_plus"))
    )
    return frame.loc[:, ~frame.columns.str.contains("^unnamed")]


def read_transactional(path: str | Path, sheet: str) -> pd.DataFrame:
    """Read a transactional sheet and trim the trailing blank rows.

    The existing ETL trims with iloc[:valid_rows] where valid_rows is a *count* of non-null
    first-column values. That is only correct while the blanks are all at the bottom -- it
    holds on 2026-07 (802 real rows then 1033 blanks) but a blank row in the middle would
    silently keep a blank and drop a real one. Masking on the column itself is equivalent
    when the blanks are trailing and correct when they are not.
    """
    frame = normalise_columns(pd.read_excel(path, dtype="str", sheet_name=sheet))
    first_column = frame.columns[0]
    mask = frame[first_column].notna()
    if not mask.iloc[: int(mask.sum())].all():
        log.warning("%s: blank rows are not contiguous at the bottom", sheet)
    return frame[mask]


def column_total(frame: pd.DataFrame, column: str) -> float | None:
    """Sum a string-typed money column, treating blanks as zero."""
    if column not in frame.columns:
        log.warning("column %r not present, cross-check skipped", column)
        return None
    return float(pd.to_numeric(frame[column], errors="coerce").fillna(0).sum())


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------


def run_tie_outs(frame: pd.DataFrame, tolerance: float) -> pd.DataFrame:
    """Check the settlement statement against itself. One row per check.

    Checks are skipped rather than failed where an input is absent on that basis -- a
    missing gross E is the template's design, not a break.
    """
    results: list[dict] = []

    def check(name: str, scope: str, expected: float, actual: float) -> None:
        expected, actual = float(expected), float(actual)
        difference = actual - expected
        results.append(
            {
                "check": name, "scope": scope,
                "expected": expected, "actual": actual, "diff": difference,
                "passed": abs(difference) <= tolerance,
            }
        )

    # Cohort columns sum to the Total column.
    for (row_ord, basis), group in frame.groupby(["row_ord", "basis"]):
        total = group.loc[group.cohort == "TOTAL", "amount"]
        parts = group.loc[group.cohort != "TOTAL", "amount"]
        if total.empty or parts.empty:
            continue
        check("cohorts_sum_to_total", f"row {row_ord} {basis}", total.iat[0], parts.sum())

    # net = gross x quota share.
    wide = frame.pivot_table(
        index=["row_ord", "cohort"], columns="basis",
        values=["amount", "quota_share"], aggfunc="first",
    )
    for (row_ord, cohort), record in wide.iterrows():
        gross = record.get(("amount", "gross"))
        net = record.get(("amount", "net"))
        quota_share = record.get(("quota_share", "net"))
        if pd.isna(gross) or pd.isna(net) or pd.isna(quota_share):
            continue
        check("net_equals_gross_x_qs", f"row {row_ord} {cohort}", gross * quota_share, net)

    # Headline settlement identities.
    coded = frame[frame.line_code.notna()]
    lookup = {(r.line_code, r.basis, r.cohort): r.amount for r in coded.itertuples()}
    for basis in sorted(frame.basis.unique()):
        for cohort in sorted(frame.cohort.unique()):
            scope = f"{basis} {cohort}"
            v = {code: lookup.get((code, basis, cohort)) for code in "ABCDEFGHIJK"}
            if all(v[c] is not None for c in "ABCD"):
                check("D = A - B - C", scope, v["A"] - v["B"] - v["C"], v["D"])
            if all(v[c] is not None for c in "DEF"):
                check("F = D - E", scope, v["D"] - v["E"], v["F"])
            if all(v[c] is not None for c in "FGH"):
                check("H = F + G", scope, v["F"] + v["G"], v["H"])
            if all(v[c] is not None for c in "HIJK"):
                check("K = H - I + J", scope, v["H"] - v["I"] + v["J"], v["K"])

    # Section 4 / 5 roll-forwards.
    for section in ROLLFORWARD_SECTIONS:
        block = frame[frame.section == section]
        if block.empty:
            continue

        anchors: dict[str, int | None] = {}
        for name, prefix in ROLLFORWARD_LABELS.items():
            hits = block[block.line_item.str.startswith(prefix, na=False)]
            anchors[name] = int(hits.row_ord.min()) if not hits.empty else None
        if any(value is None for value in anchors.values()):
            log.warning("section %d: roll-forward anchors not found, skipping those checks", section)
            continue

        for basis in sorted(block.basis.unique()):
            for cohort in sorted(block.cohort.unique()):
                amounts = (
                    block[(block.basis == basis) & (block.cohort == cohort)]
                    .set_index("row_ord")["amount"]
                )
                begin, change = amounts.get(anchors["begin"]), amounts.get(anchors["change"])
                end = amounts.get(anchors["end"])
                components = amounts[
                    (amounts.index > anchors["begin"]) & (amounts.index < anchors["change"])
                ]
                scope = f"s{section} {basis} {cohort}"
                if change is not None and not components.empty:
                    check("change = sum(components)", scope, components.sum(), change)
                if None not in (begin, change, end):
                    check("end = begin + change", scope, begin + change, end)

    return pd.DataFrame(results)


def run_cross_checks(frame: pd.DataFrame, path: str | Path, tolerance: float) -> pd.DataFrame:
    """Reconcile the settlement against the seriatim and transactional sheets.

    Confirmed to the cent on 2026-07. Note the sign convention on surrender_charge: it is
    stored negative on both the Withdrawals sheet and the settlement, so claims are
    withdrawal_amount + surrender_charge, not minus.
    """
    results: list[dict] = []

    def check(name: str, source: str, expected, actual) -> None:
        if expected is None or actual is None:
            results.append({"check": name, "source": source, "settlement": None,
                            "source_value": None, "diff": None, "passed": None})
            return
        difference = float(actual) - float(expected)
        results.append(
            {
                "check": name, "source": source,
                "settlement": float(expected), "source_value": float(actual),
                "diff": difference, "passed": abs(difference) <= tolerance,
            }
        )

    def line(code: str, basis: str = "gross", cohort: str = "TOTAL"):
        hit = frame[(frame.line_code == code) & (frame.basis == basis) & (frame.cohort == cohort)]
        return None if hit.empty else float(hit.amount.iat[0])

    def rollforward(section: int, prefix: str, basis: str = "gross", cohort: str = "TOTAL"):
        block = frame[
            (frame.section == section) & (frame.basis == basis) & (frame.cohort == cohort)
            & frame.line_item.str.startswith(prefix, na=False)
        ]
        return None if block.empty else float(block.sort_values("row_ord").amount.iat[0])

    try:
        seriatim = read_transactional(path, "Seriatim")
        premium = read_transactional(path, "Premiums")
        withdrawals = read_transactional(path, "Withdrawals")
    except Exception as exc:
        log.warning("could not read transactional sheets, cross-checks skipped: %s", exc)
        return pd.DataFrame(results)

    log.info("cross-check sources: seriatim %d, premiums %d, withdrawals %d rows",
             len(seriatim), len(premium), len(withdrawals))

    check("A total premium", "Premiums.total_premium",
          line("A"), column_total(premium, "total_premium"))

    withdrawal_amount = column_total(withdrawals, "withdrawal_amount")
    surrender_charge = column_total(withdrawals, "surrender_charge")
    claims = (None if None in (withdrawal_amount, surrender_charge)
              else withdrawal_amount + surrender_charge)
    check("B total claims", "Withdrawals.withdrawal_amount + surrender_charge",
          line("B"), claims)

    check("s5 account value begin", "Seriatim.bom_fund_value",
          rollforward(5, "Values at the Beginning"), column_total(seriatim, "bom_fund_value"))
    check("s5 account value end", "Seriatim.eom_fund_value",
          rollforward(5, "Values at the End"), column_total(seriatim, "eom_fund_value"))

    return pd.DataFrame(results)


def report_validation(tie_outs: pd.DataFrame, cross_checks: pd.DataFrame,
                      fail: bool, set_month: str) -> None:
    """Log both check tables. Raises ValidationError on failure when fail=True."""
    tie_failures = tie_outs[~tie_outs.passed]
    log.info(
        "internal tie-outs %s: %d checks, %d passed, %d failed",
        set_month, len(tie_outs), int(tie_outs.passed.sum()), len(tie_failures),
    )
    for name, group in tie_outs.groupby("check"):
        bad = int((~group.passed).sum())
        log.log(logging.ERROR if bad else logging.DEBUG,
                "  %-26s %4d checks  %d failed", name, len(group), bad)
    for record in tie_failures.itertuples():
        log.error("  FAIL %s [%s] expected %.4f actual %.4f diff %+.4f",
                  record.check, record.scope, record.expected, record.actual, record.diff)

    if cross_checks.empty:
        cross_failures = cross_checks
    else:
        # Compare with == rather than `is`: when every check passes the column comes back as
        # numpy bool, and numpy.bool_(True) is not the True singleton.
        cross_failures = cross_checks[cross_checks.passed == False]  # noqa: E712
        log.info(
            "cross-sheet checks %s: %d checks, %d passed, %d failed, %d skipped",
            set_month, len(cross_checks), int((cross_checks.passed == True).sum()),  # noqa: E712
            len(cross_failures), int(cross_checks.passed.isna().sum()),
        )
        for record in cross_checks.itertuples():
            passed = None if pd.isna(record.passed) else bool(record.passed)
            if passed is None:
                log.warning("  SKIP %-24s (%s unavailable)", record.check, record.source)
            elif passed:
                log.debug("  OK   %-24s %s = %s", record.check, record.source,
                          f"{record.source_value:,.2f}")
            else:
                log.error("  FAIL %-24s settlement %.2f vs %s %.2f (diff %+.2f)",
                          record.check, record.settlement, record.source,
                          record.source_value, record.diff)

    total_failures = len(tie_failures) + len(cross_failures)
    if total_failures and fail:
        raise ValidationError(
            f"{total_failures} validation failure(s) for {set_month} -- not loading. "
            "The template may have changed."
        )


# ---------------------------------------------------------------------------
# BigQuery load
# ---------------------------------------------------------------------------


def resolve_credentials(credentials):
    """Coerce whatever the caller passed into a google-auth Credentials object.

    Accepts an already-built Credentials object, a path to a service account JSON, the
    parsed contents of one as a dict, or None for application-default credentials.

    The path form exists because KSKJ etl.ipynb keeps its service account as a path in
    CREDS, so `credentials=CREDS` is the natural thing to write. Passing that string
    straight to bigquery.Client raises a fairly opaque ValueError about
    google-auth-library-python, so it is worth converting here rather than making every
    caller remember to.
    """
    if credentials is None:
        return None

    if isinstance(credentials, (str, Path)):
        # Check the path before importing, so a typo'd path reports itself rather than
        # surfacing as an import error.
        path = Path(credentials)
        if not path.exists():
            raise FileNotFoundError(f"service account JSON not found: {path}")

        from google.oauth2 import service_account

        log.debug("building credentials from %s", path)
        return service_account.Credentials.from_service_account_file(str(path))

    if isinstance(credentials, dict):
        from google.oauth2 import service_account

        return service_account.Credentials.from_service_account_info(credentials)

    # Anything else is assumed to be a google-auth Credentials object already. Validate it
    # here so the error names the argument instead of surfacing from deep inside the client.
    try:
        from google.auth.credentials import Credentials
    except ImportError:  # google-auth absent; let the client complain
        return credentials

    if not isinstance(credentials, Credentials):
        raise TypeError(
            f"credentials must be a google-auth Credentials object, a path to a service "
            f"account JSON, a dict, or None -- got {type(credentials).__name__}"
        )
    return credentials


def to_bq_numeric(value):
    """Convert a float to an exact Decimal suitable for a BigQuery NUMERIC column.

    BigQuery NUMERIC is DECIMAL(38, 9). pandas-gbq loads via Parquet, and pyarrow will not
    convert a float64 to a 9-scale decimal unless the float's *exact* binary value fits in
    nine decimal places -- otherwise it raises
    'ArrowInvalid: Rescaling Decimal128 value would cause data loss'.

    Almost no money value survives that test. 1408390.9 is exactly
    1408390.8999999999068677425384521484375 as a double, because tenths are not
    representable in binary; 332 of the 1428 amounts on the 2026-07 statement fail.

    Building the Decimal from str(value) rather than from the float is what fixes it: str()
    gives the shortest repr that round-trips (1408390.9), so quantizing to nine places is
    exact and pyarrow has nothing to rescale. Decimal(value) would instead capture the full
    binary artefact and fail all over again.
    """
    if value is None or pd.isna(value):
        return None
    return Decimal(str(value)).quantize(NUMERIC_QUANTUM, rounding=ROUND_HALF_UP)


def prepare_for_bigquery(frame: pd.DataFrame) -> pd.DataFrame:
    """Cast dtypes so the frame lands with the intended BigQuery types.

    report_date becomes the dbdate extension dtype so it loads as a true DATE rather than a
    TIMESTAMP at midnight. db-dtypes ships as a pandas-gbq dependency.

    amount becomes exact Decimals for the NUMERIC column -- see to_bq_numeric. quota_share
    stays float64 because its column is declared FLOAT64.
    """
    prepared = frame.copy()
    prepared["report_date"] = prepared["report_date"].astype("dbdate")
    prepared["section"] = prepared["section"].astype("Int64")
    prepared["row_ord"] = prepared["row_ord"].astype("Int64")
    prepared["amount"] = prepared["amount"].map(to_bq_numeric)
    return prepared[COLUMN_ORDER]


def load_to_bigquery(frame: pd.DataFrame, project: str, dataset: str, table: str,
                     cedent: str, credentials=None) -> None:
    """Delete the month's existing rows, then append.

    pandas_gbq.to_gbq has no upsert or DML, so idempotency comes from an explicit DELETE
    scoped to this set_month followed by an append.

    Two things make that safe. to_gbq uses a load job rather than the streaming API, so the
    freshly written rows are not held in a streaming buffer and next month's DELETE will see
    them. And the DELETE and append are separate operations, not a transaction -- if the
    append fails the month is left empty rather than duplicated, which is the right failure
    direction for a reconciliation table. Absent data is obvious; duplicated data is
    silently wrong. Re-run to fix.
    """
    import pandas_gbq
    from google.cloud import bigquery

    fqn = f"{project}.{dataset}.{table}"
    set_month = frame.set_month.iat[0]

    credentials = resolve_credentials(credentials)
    client = bigquery.Client(project=project, credentials=credentials)

    client.query(DDL_TEMPLATE.format(fqn=fqn)).result()
    log.info("table ready: %s", fqn)

    delete_job = client.query(
        DELETE_TEMPLATE.format(fqn=fqn),
        job_config=bigquery.QueryJobConfig(
            query_parameters=[
                bigquery.ScalarQueryParameter("set_month", "STRING", set_month),
                bigquery.ScalarQueryParameter("cedent", "STRING", cedent),
            ]
        ),
    )
    delete_job.result()
    log.info("deleted %s existing row(s) for %s", delete_job.num_dml_affected_rows, set_month)

    pandas_gbq.to_gbq(
        frame,
        destination_table=f"{dataset}.{table}",
        project_id=project,
        if_exists="append",
        table_schema=BQ_SCHEMA,
        credentials=credentials,
        progress_bar=False,
    )
    log.info("appended %d rows for %s to %s", len(frame), set_month, fqn)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def load_settlement(
    path: str | Path,
    set_month: str | None = None,
    project: str = PROJECT,
    dataset: str = DATASET,
    table: str = TABLE,
    cedent: str = CEDENT,
    treaty: str = TREATY,
    credentials=None,
    tolerance: float = 0.01,
    fail_on_validation: bool = True,
    cross_check: bool = True,
    dry_run: bool = False,
) -> pd.DataFrame:
    """Extract, validate and load one monthly workbook's Settlement tab.

    This is the drop-in for KSKJ etl.ipynb. Pass set_month to assert that the workbook's own
    period matches what the rest of the notebook derived from the Seriatim sheet -- a
    mismatch means the wrong file, which is worth catching loudly.

    credentials accepts a google-auth Credentials object, a path to a service account JSON
    (so `credentials=CREDS` works directly), a dict, or None for application-default.

    Returns the extracted frame.
    """
    # Resolve credentials before the slow work, so a bad path fails in milliseconds rather
    # than after parsing and validating the whole workbook.
    if not dry_run:
        credentials = resolve_credentials(credentials)

    frame = extract_settlement(path, cedent=cedent, treaty=treaty)
    workbook_month = frame.set_month.iat[0]

    if set_month is not None and str(set_month) != workbook_month:
        raise ValueError(
            f"set_month mismatch: pipeline says {set_month}, "
            f"{Path(path).name} Settlement!{PERIOD_CELL} says {workbook_month}"
        )

    tie_outs = run_tie_outs(frame, tolerance=tolerance)
    cross_checks = (run_cross_checks(frame, path, tolerance=tolerance)
                    if cross_check else pd.DataFrame())
    report_validation(tie_outs, cross_checks, fail=fail_on_validation, set_month=workbook_month)

    if dry_run:
        log.info("dry run -- would load %d rows for %s to %s.%s.%s",
                 len(frame), workbook_month, project, dataset, table)
        return frame

    load_to_bigquery(
        prepare_for_bigquery(frame),
        project=project, dataset=dataset, table=table,
        cedent=cedent, credentials=credentials,
    )
    return frame


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def resolve_paths(patterns: list[str]) -> list[Path]:
    """Expand glob patterns, keeping literal paths that exist. Skips Excel lock files."""
    paths: list[Path] = []
    for pattern in patterns:
        matches = [Path(p) for p in sorted(glob.glob(pattern))]
        if not matches:
            candidate = Path(pattern)
            if not candidate.exists():
                raise FileNotFoundError(f"no file matches {pattern!r}")
            matches = [candidate]
        paths.extend(p for p in matches if not p.name.startswith("~$"))
    if not paths:
        raise FileNotFoundError(f"no workbooks matched {patterns}")
    return paths


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Load the Settlement tab of the KSKJ monthly report into BigQuery.",
    )
    parser.add_argument("workbooks", nargs="+",
                        help="path(s) or glob pattern(s) to monthly .xlsx report(s)")
    parser.add_argument("--creds", default=os.environ.get("GOOGLE_APPLICATION_CREDENTIALS", CREDS),
                        help=f"service account JSON [default: {CREDS}]")
    parser.add_argument("--project", default=PROJECT, help=f"GCP project [default: {PROJECT}]")
    parser.add_argument("--dataset", default=DATASET, help=f"dataset [default: {DATASET}]")
    parser.add_argument("--table", default=TABLE, help=f"table [default: {TABLE}]")
    parser.add_argument("--cedent", default=CEDENT, help="cedent code stored on each row")
    parser.add_argument("--treaty", default=TREATY, help="treaty code stored on each row")
    parser.add_argument("--set-month", help="assert the workbook's period equals this YYYYMM")
    parser.add_argument("--tolerance", type=float, default=0.01,
                        help="check tolerance in dollars (default: 0.01)")
    parser.add_argument("--dry-run", action="store_true",
                        help="extract and validate only; no BigQuery writes")
    parser.add_argument("--no-cross-check", action="store_true",
                        help="skip reconciliation against the seriatim/transactional sheets")
    parser.add_argument("--no-fail-on-validation", action="store_true",
                        help="log check failures but load anyway (not recommended)")
    parser.add_argument("--csv", metavar="PATH", help="also write the extracted rows to this CSV")
    parser.add_argument("-v", "--verbose", action="store_true", help="debug logging")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
    )
    warnings.filterwarnings("ignore", message=".*Workbook contains no default style.*")

    credentials = None
    if not args.dry_run:
        try:
            credentials = resolve_credentials(args.creds)
        except (FileNotFoundError, TypeError, ValueError) as exc:
            log.error("%s", exc)
            return 2
        log.info("authenticated as %s", getattr(credentials, "service_account_email", "?"))

    try:
        paths = resolve_paths(args.workbooks)
    except FileNotFoundError as exc:
        log.error("%s", exc)
        return 2

    log.info("%d workbook(s) to process", len(paths))

    frames: list[pd.DataFrame] = []
    failed: list[tuple[Path, Exception]] = []

    for path in paths:
        try:
            frames.append(
                load_settlement(
                    path,
                    set_month=args.set_month,
                    project=args.project, dataset=args.dataset, table=args.table,
                    cedent=args.cedent, treaty=args.treaty,
                    credentials=credentials,
                    tolerance=args.tolerance,
                    fail_on_validation=not args.no_fail_on_validation,
                    cross_check=not args.no_cross_check,
                    dry_run=args.dry_run,
                )
            )
        except Exception as exc:  # keep going; report at the end
            log.error("%s: %s: %s", path.name, type(exc).__name__, exc)
            failed.append((path, exc))

    if args.csv and frames:
        combined = pd.concat(frames, ignore_index=True)
        combined.to_csv(args.csv, index=False)
        log.info("wrote %d rows to %s", len(combined), args.csv)

    log.info("done: %d succeeded, %d failed", len(frames), len(failed))
    for path, exc in failed:
        log.error("  %s -- %s", path.name, exc)

    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
